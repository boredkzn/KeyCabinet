#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Генерация отчёта ДЗ5 (ГОСТ-подобное оформление). Зарипов / KeyCabinet + Supabase."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parents[1] / "Отчет_ДЗ5_Зарипов.docx"


def set_run_font(run, name="Times New Roman", size=14, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def setup_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.different_first_page_header_footer = True

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(14)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, center in (("Heading 1", True), ("Heading 2", False)):
        st = doc.styles[style_name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(14)
        st.font.bold = True
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        )


def add_page_numbers(doc: Document):
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.clear()
    run = p.add_run()
    set_run_font(run, size=12)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    fp = first_footer.paragraphs[0] if first_footer.paragraphs else first_footer.add_paragraph()
    fp.clear()


def add_center(doc, text, bold=False, size=14, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_body(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.clear()
    run = p.add_run(text)
    set_run_font(run)
    return p


def build_title(doc: Document):
    for _ in range(2):
        add_center(doc, "")
    add_center(doc, "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ", bold=True, size=12)
    add_center(doc, "РОССИЙСКОЙ ФЕДЕРАЦИИ", bold=True, size=12)
    add_center(doc, "")
    add_center(doc, "Федеральное государственное бюджетное образовательное", size=12)
    add_center(doc, "учреждение высшего образования", size=12)
    add_center(doc, "")
    add_center(doc, "ОТЧЁТ", bold=True, size=16, space_after=6)
    add_center(doc, "по дисциплине «Методы искусственного интеллекта»", bold=True)
    add_center(doc, "Домашнее задание №5", bold=True, space_after=12)
    add_center(doc, "«Развертывание Backend и интеграция с Frontend»", space_after=6)
    add_center(doc, "Проект KeyCabinet · Backend: Supabase (BaaS)", space_after=24)
    add_center(doc, "Выполнил: студент Зарипов Шамиль Дамирович", space_after=6)
    add_center(doc, "Группа: 4481", space_after=24)
    add_center(doc, "2026", space_after=0)
    doc.add_page_break()


def build_body(doc: Document):
    doc.add_heading("1. Цель работы", level=1)
    add_body(
        doc,
        "Спроектировать схему базы данных приложения KeyCabinet, развернуть backend "
        "на платформе Supabase (Backend as a Service), настроить аутентификацию и "
        "политики Row Level Security, интегрировать Frontend с PostgREST API и "
        "реализовать обработку сетевых ошибок (401, 403, 500).",
    )

    doc.add_heading("2. Шаг 1. Проектирование базы данных", level=1)
    add_body(
        doc,
        "С помощью AI-агента (Cursor) по требованиям KeyCabinet сгенерирована SQL-схема. "
        "Сущности: office_keys (ключи кабинетов), checkouts (история выдач), "
        "profiles (профили пользователей Auth).",
    )
    add_bullet(doc, "office_keys: id, room, label, floor, status, holder, issued_at")
    add_bullet(doc, "checkouts: id, key_id, holder_name, issued_at, returned_at, issued_by")
    add_bullet(doc, "profiles: id → auth.users, display_name, role")
    add_body(
        doc,
        "Миграция сохранена в файле supabase/migrations/001_init.sql и включает seed "
        "из пяти демо-ключей, согласованных с Frontend ДЗ4.",
    )

    doc.add_heading("3. Шаг 2. Выбор инфраструктурного решения", level=1)
    add_body(
        doc,
        "Выбран вариант A — Supabase (BaaS). Self-hosted PostgreSQL в Docker на VPS "
        "отклонён: для учебного проекта избыточен (аренда сервера, Docker, TLS, "
        "ручная настройка Auth и CORS). Supabase даёт Auth, PostgREST, RLS и логи "
        "из коробки на Free tier, что соответствует предпочтению не покупать VPS.",
    )

    doc.add_heading("4. Шаг 3. Развёртывание базы данных", level=1)
    add_body(
        doc,
        "Подготовлены SQL-миграция и инструкция создания проекта в Dashboard. "
        "Создание облачного проекта требует входа студента в supabase.com: "
        "New project → SQL Editor → выполнить 001_init.sql → включить Email Auth. "
        "До заполнения .env.local Frontend работает в demo-режиме.",
    )

    doc.add_heading("5. Шаг 4. Создание API", level=1)
    add_body(
        doc,
        "Используется автоматический REST API PostgREST. Клиентский слой "
        "src/api/keysApi.ts реализует чтение списка ключей (GET), выдачу "
        "(PATCH office_keys + POST checkouts) и возврат (PATCH office_keys + "
        "закрытие checkout). Минимум три CRUD-операции закрыты.",
    )

    doc.add_heading("6. Шаг 5. Настройка безопасности", level=1)
    add_bullet(doc, "Supabase Auth: регистрация и вход email/password")
    add_bullet(doc, "RLS на profiles, office_keys, checkouts (только authenticated)")
    add_bullet(doc, "Секреты только в .env.local; в репозитории — .env.example")
    add_bullet(doc, "CORS: стандартная конфигурация Supabase для браузерного anon key")

    doc.add_heading("7. Шаг 6. Интеграция Frontend с Backend", level=1)
    add_body(
        doc,
        "Установлен @supabase/supabase-js. Модули src/lib/supabase.ts и src/api/* "
        "заменяют чистый mock: при наличии VITE_SUPABASE_URL и VITE_SUPABASE_ANON_KEY "
        "данные идут с сервера; иначе — graceful demo-режим с баннером, чтобы "
        "npm run dev / Run ▶ в Rider продолжали работать.",
    )

    doc.add_heading("8. Шаг 7. Обработка ошибок и логирование", level=1)
    add_body(
        doc,
        "Модуль src/lib/apiErrors.ts преобразует ответы PostgREST/Auth и сетевые "
        "сбои в сообщения UI для 401, 403, 500 и offline. Логи просматриваются в "
        "Supabase Dashboard → Logs. Для AI-отладки достаточно вставить фрагмент "
        "лога и текст RLS-политики.",
    )

    doc.add_heading("9. Шаг 8. Тестирование и отладка", level=1)
    add_body(
        doc,
        "Расширен набор Vitest: утилиты, apiErrors, keysApi (demo + mock-клиент), "
        "UI App. Команда npm test выполняется без live Supabase. На момент сдачи: "
        "25 тестов, все успешно. Опциональная интеграция — после заполнения .env.local.",
    )

    doc.add_heading("10. Шаг 9. Оформление результатов", level=1)
    add_body(
        doc,
        "Подготовлены backend_documentation.md (архитектура, деплой, endpoints, "
        "процесс с AI), обновлён README.md, отчёт в формате Word, пакет сдачи "
        "на Desktop (МИИ Зарипов/5).",
    )

    doc.add_heading("11. Выводы", level=1)
    add_body(
        doc,
        "Backend KeyCabinet спроектирован и интегрирован через Supabase. "
        "Приложение сохраняет работоспособность без ключей (demo) и готово к "
        "live-режиму после создания проекта студентом. AI ускорил генерацию "
        "SQL/RLS, клиента API, тестов и документации.",
    )

    doc.add_heading("12. Приложения", level=1)
    add_bullet(doc, "Рабочая копия: ~/RiderProjects/KeyCabinet/")
    add_bullet(doc, "Миграция: supabase/migrations/001_init.sql")
    add_bullet(doc, "Документация: backend_documentation.md")
    add_bullet(doc, "Сдача: ~/Desktop/МИИ Зарипов/5/")


def main():
    doc = Document()
    setup_styles(doc)
    add_page_numbers(doc)
    build_title(doc)
    build_body(doc)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
